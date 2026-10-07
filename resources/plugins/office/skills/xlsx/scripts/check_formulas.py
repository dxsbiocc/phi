#!/usr/bin/env python3
"""Statically inspect formulas in an XLSX or XLSM workbook and emit JSON."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Iterable, Iterator

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries

ERROR_LITERAL_PATTERN = re.compile(r"#(?:REF!|NAME\?)", re.IGNORECASE)
WHOLE_COLUMN_PATTERN = re.compile(r"(?<![A-Z0-9_])\$?[A-Z]{1,3}:\$?[A-Z]{1,3}(?![A-Z0-9_])", re.IGNORECASE)
DIVIDE_ZERO_PATTERN = re.compile(
    r"/\s*(?:\(\s*)*[+-]?(?:0+(?:\.0*)?|\.0+)(?:\s*\))*\s*(?=$|[+\-*/^,;)])"
)
CELL_REFERENCE_PATTERN = re.compile(
    r"^(?:(?:'((?:[^']|'')+)'|([^!]+))!)?(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?$",
    re.IGNORECASE,
)

CellKey = tuple[str, str]
MAX_ISSUES = 400
MAX_FORMULA_CHARS = 300


def balanced_parentheses(formula: str) -> bool:
    depth = 0
    quote: str | None = None
    index = 0
    while index < len(formula):
        char = formula[index]
        if quote:
            if char == quote:
                if index + 1 < len(formula) and formula[index + 1] == quote:
                    index += 1
                else:
                    quote = None
        elif char in {'"', "'"}:
            quote = char
        elif char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
            if depth < 0:
                return False
        index += 1
    return depth == 0 and quote is None


def referenced_sheets(formula: str) -> set[str]:
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return set()
    sheets: set[str] = set()
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        match = CELL_REFERENCE_PATTERN.fullmatch(token.value.strip())
        if not match:
            continue
        quoted_sheet, plain_sheet, _, _ = match.groups()
        if quoted_sheet is not None:
            sheets.add(quoted_sheet.replace("''", "'"))
        elif plain_sheet is not None:
            sheets.add(plain_sheet.strip())
    return sheets


def range_dependencies(value: str, current_sheet: str, formula_cells: set[CellKey]) -> set[CellKey]:
    match = CELL_REFERENCE_PATTERN.fullmatch(value.strip())
    if not match:
        return set()
    quoted_sheet, plain_sheet, start, end = match.groups()
    sheet = quoted_sheet.replace("''", "'") if quoted_sheet is not None else plain_sheet
    sheet = sheet.strip() if sheet else current_sheet
    start = start.replace("$", "").upper()
    end = (end or start).replace("$", "").upper()
    try:
        min_col, min_row, max_col, max_row = range_boundaries(f"{start}:{end}")
    except ValueError:
        return set()
    return {
        key
        for key in formula_cells
        if key[0].casefold() == sheet.casefold()
        and min_col <= _column_index(key[1]) <= max_col
        and min_row <= _row_index(key[1]) <= max_row
    }


def _column_index(coordinate: str) -> int:
    letters = re.match(r"[A-Z]+", coordinate.upper())
    if not letters:
        return 0
    value = 0
    for char in letters.group(0):
        value = value * 26 + ord(char) - ord("A") + 1
    return value


def _row_index(coordinate: str) -> int:
    digits = re.search(r"\d+$", coordinate)
    return int(digits.group(0)) if digits else 0


def formula_dependencies(formula: str, sheet: str, formula_cells: set[CellKey]) -> set[CellKey]:
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return set()
    dependencies: set[CellKey] = set()
    for token in tokens:
        if token.type == "OPERAND" and token.subtype == "RANGE":
            dependencies.update(range_dependencies(token.value, sheet, formula_cells))
    return dependencies


def cycle_nodes(graph: dict[CellKey, set[CellKey]]) -> set[CellKey]:
    state: dict[CellKey, int] = {}
    found: set[CellKey] = set()
    for start in graph:
        if state.get(start, 0) != 0:
            continue
        path = [start]
        positions = {start: 0}
        state[start] = 1
        stack: list[tuple[CellKey, Iterator[CellKey]]] = [
            (start, iter(sorted(graph.get(start, set()))))
        ]
        while stack:
            node, targets = stack[-1]
            try:
                target = next(targets)
            except StopIteration:
                stack.pop()
                positions.pop(node)
                path.pop()
                state[node] = 2
                continue
            target_state = state.get(target, 0)
            if target_state == 0:
                state[target] = 1
                positions[target] = len(path)
                path.append(target)
                stack.append((target, iter(sorted(graph.get(target, set())))))
            elif target_state == 1:
                found.update(path[positions[target] :])
    return found


def issue(sheet: str, cell: str, code: str, message: str, formula: str) -> dict[str, str]:
    return {
        "sheet": sheet,
        "cell": cell,
        "code": code,
        "message": message,
        "formula": formula[:MAX_FORMULA_CHARS],
    }


def local_issues(sheet: str, cell: str, formula: str, sheetnames: set[str]) -> Iterable[dict[str, str]]:
    if ERROR_LITERAL_PATTERN.search(formula):
        yield issue(sheet, cell, "error_literal", "Formula contains #REF! or #NAME?.", formula)
    if not balanced_parentheses(formula):
        yield issue(sheet, cell, "unbalanced_parentheses", "Formula parentheses are unbalanced.", formula)
    if WHOLE_COLUMN_PATTERN.search(formula):
        yield issue(sheet, cell, "whole_column_reference", "Formula uses a whole-column reference.", formula)
    if DIVIDE_ZERO_PATTERN.search(formula):
        yield issue(sheet, cell, "division_by_literal_zero", "Formula divides by the literal value zero.", formula)
    for referenced in sorted(referenced_sheets(formula)):
        if referenced.casefold() not in sheetnames:
            yield issue(sheet, cell, "missing_sheet", f"Referenced worksheet does not exist: {referenced}", formula)


def inspect_workbook(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=False, read_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    formulas: dict[CellKey, str] = {}
    issues: list[dict[str, str]] = []
    sheetnames = {name.casefold() for name in workbook.sheetnames}
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f" and isinstance(cell.value, str):
                    formula = cell.value if cell.value.startswith("=") else f"={cell.value}"
                    key = (worksheet.title, cell.coordinate)
                    formulas[key] = formula
                    issues.extend(local_issues(*key, formula, sheetnames))

    formula_cells = set(formulas)
    graph = {
        key: formula_dependencies(formula, key[0], formula_cells)
        for key, formula in formulas.items()
    }
    for sheet, cell in sorted(cycle_nodes(graph)):
        issues.append(issue(sheet, cell, "circular_reference", "Formula participates in an obvious circular reference.", formulas[(sheet, cell)]))

    issues.sort(key=lambda item: (item["sheet"], item["cell"], item["code"]))
    issue_count = len(issues)
    returned_issues = issues[:MAX_ISSUES]
    return {
        "file": path.name,
        "sheetCount": len(workbook.sheetnames),
        "formulaCount": len(formulas),
        "ok": issue_count == 0,
        "issueCount": issue_count,
        "issuesTruncated": issue_count > len(returned_issues),
        "issues": returned_issues,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        result = inspect_workbook(args.input)
    except Exception as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False))
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
